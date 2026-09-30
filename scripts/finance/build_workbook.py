"""
Builds gylio-financial-model.xlsx: the same model as server/billing/financeModel.js,
written as live spreadsheet formulas so every number moves when an input changes.

    python3 scripts/finance/build_workbook.py <output.xlsx>

Yellow cells with blue text are inputs. Black text is a formula on its own sheet;
green text reaches into another sheet. Inputs are read from the JS assumptions via
`node scripts/finance/dump_assumptions.cjs`, so the workbook and the code start
from identical numbers.
"""
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parents[2]
MONTHS = 36
FIRST_ROW = 20  # first month row on the Cash flow sheet

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1D6B57")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="D5DBD6"))
USD = '"$"#,##0.00;("$"#,##0.00);"-"'
USD0 = '"$"#,##0;("$"#,##0);"-"'
PEN = '"S/ "#,##0.00;("S/ "#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
MULT = '0.0"x"'


def assumptions():
    out = subprocess.run(["node", "scripts/finance/dump_assumptions.cjs"], cwd=ROOT, check=True,
                         capture_output=True, text=True).stdout
    return json.loads(out)


def name(wb, label, ref):
    wb.defined_names[label] = DefinedName(label, attr_text=ref)


def header(ws, row, labels, start_col=1):
    for i, label in enumerate(labels):
        cell = ws.cell(row=row, column=start_col + i, value=label)
        cell.font = HEAD
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def input_cell(cell, value, fmt=None, note=None):
    cell.value = value
    cell.font = BLUE
    cell.fill = INPUT_FILL
    if fmt:
        cell.number_format = fmt
    if note:
        cell.comment = Comment(note, "Gylio model")


# ── Read me ──────────────────────────────────────────────────────────────────
def build_readme(ws):
    ws["A1"] = "Gylio Pro — the money model"
    ws["A1"].font = TITLE
    ws["A2"] = ("Private. Same model as server/billing/financeModel.js, as live formulas. "
                "Money is in US dollars unless a label says soles (S/).")
    rows = [
        ("How to use this", ""),
        ("Yellow cells, blue text", "The only cells to change. Every other number is a formula and moves on its own."),
        ("Black text", "A formula that works inside its own sheet."),
        ("Green text", "A formula that reaches into another sheet."),
        ("", ""),
        ("The sheets", ""),
        ("Assumptions", "Every input: prices, exchange rate, taxes, Paddle's fee, demand for three scenarios, fixed costs, your hours."),
        ("Unit economics", "What one charge leaves you after IGV/sales tax, Paddle's 5% + $0.50 and the soles-to-dollars spread."),
        ("Break-even", "Monthly fixed costs and how many paying users cover them; how the contador's fee moves that number."),
        ("Price equilibrium", "The soles monthly price against how much demand a price rise costs (linear demand, four elasticities)."),
        ("Cash flow", "36 months, three scenarios side by side: NPV, payback month, return on your hours, users and MRR."),
        ("Sensitivity", "NPV against the discount rate, and against trial conversion and churn together."),
        ("Sensitivity engine", "The arithmetic behind the conversion × churn table. You never need to open it; nothing is hidden."),
        ("", ""),
        ("What this model does not do", ""),
        ("No paid acquisition", "No ad spend or cost per install. If you buy traffic, add it to the fixed costs."),
        ("No refunds or chargebacks", "Paddle keeps its fee on refunds; at this size they are noise. Add a point to the fee if you want them."),
        ("Income tax is one flat rate", "RMT is 10% up to 15 UIT (S/ 82,500 in 2026) of annual profit, then 29.5%. The 1% monthly advance is timing, not cost."),
        ("The one real judgement call", "Demand's response to price (elasticity). Nobody knows it for Gylio yet, so the price sheet shows a band."),
        ("Not yet measured", "Trial start rate, trial-to-paid and churn are benchmarks, not Gylio data. Replace them after the beta."),
    ]
    for i, (a, b) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=a).font = BOLD if a and not b else Font(name=FONT)
        ws.cell(row=i, column=2, value=b).font = Font(name=FONT)
        ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110


# ── Assumptions ──────────────────────────────────────────────────────────────
SCENARIO_ROWS = [
    # key, label, format, note
    ("signups0", "Sign-ups in month 1", NUM, "Assumption. Organic and content; no paid acquisition."),
    ("signupGrowth", "Sign-up growth in month 2", PCT, "Assumption. Fades each month (next row)."),
    ("growthFade", "Growth fade per month", '0.00', "Growth in month n is growth × fade^(n−2). 8% compounded for 3 years is 15×; fading avoids that."),
    ("betaMonths", "Gifted beta months (nobody charged)", NUM, "Run the beta on gifted Pro before registering a RUC."),
    ("trialStartRate", "Share of new sign-ups who start the trial", PCT, "Adapty: 31–65% depending on paywall timing."),
    ("trialToPaid", "Trial to paid (no card)", PCT, "Recurly: ~12% without a card, ~40% with a card up front."),
    ("freemiumMonthly", "Free users who convert each month", PCT, "Freemium median ~2.1% lifetime by day 35."),
    ("annualShare", "Payers choosing yearly", PCT, "Productivity apps ~77% monthly; the paywall defaults to yearly."),
    ("monthlyChurn", "Monthly plan churn", PCT, "Consumer subscriptions 7–10%."),
    ("annualRenewal", "Yearly plans that renew", PCT, "About half; 35% of yearly cancellations happen in month 1."),
    ("freeRetention", "Free users still active next month", PCT, "Assumption."),
    ("peruShare", "Payers in Peru (paying in soles)", PCT, "Spanish-first launch."),
]

COMMON_ROWS = [
    ("penPerUsd", "Soles per US dollar", '0.00', "SBS rate, late 2026 (assumption)."),
    ("priceMonthlyUSD", "Pro monthly, USD (tax included)", USD, "server/billing/plans.js"),
    ("priceYearlyUSD", "Pro yearly, USD (tax included)", USD, "server/billing/plans.js"),
    ("priceMonthlyPEN", "Pro monthly, soles (tax included)", PEN, "server/billing/plans.js"),
    ("priceYearlyPEN", "Pro yearly, soles (tax included)", PEN, "server/billing/plans.js"),
    ("feePct", "Paddle fee, % of gross", PCT, "paddle.com/pricing: 5% + $0.50, merchant of record."),
    ("feeFixed", "Paddle fee, fixed per charge", USD, "paddle.com/pricing"),
    ("fxSpread", "Soles → dollars conversion spread", PCT, "Up to ~1.5% (secondary source). Check on your first payout."),
    ("taxPeru", "IGV on Peruvian buyers", PCT, "18%; the price shown must include it."),
    ("taxIntl", "Blended sales tax / VAT elsewhere", PCT, "Assumption: many US states exempt SaaS, EU ~20%."),
    ("mau1", "Infra anchor 1: active users", NUM, "Hosting research: Cloud Run + Atlas + Cloudflare Pages."),
    ("cost1", "Infra anchor 1: $ / month", USD, ""),
    ("mau2", "Infra anchor 2: active users", NUM, ""),
    ("cost2", "Infra anchor 2: $ / month", USD, "Atlas Flex from here."),
    ("mau3", "Infra anchor 3: active users", NUM, ""),
    ("cost3", "Infra anchor 3: $ / month", USD, "Above this, scales linearly."),
    ("contadorPen", "Contador, S/ a month (after the RUC)", PEN, "RMT entry price. Get three quotes: every S/ 100 is ~8 paying users."),
    ("invoicingPen", "Electronic invoicing, S/ a month", PEN, "0 with Paddle only (one monthly export factura on free SEE-SOL). S/ 70 (NubeFact) if Mercado Pago passes are on."),
    ("payoutFeeMonthly", "Payout wire fee, $ a month", USD, "Up to $15 when Paddle pays out."),
    ("adsStartMonth", "Ads start after month", NUM, "AdSense approval takes time; house ads earn nothing."),
    ("adPageviewsPerFreeUser", "Ad pageviews per free user a month", NUM, "Ads only on Settings and QA."),
    ("adRpmUsd", "Blended ad RPM, $ per 1,000 views", USD, "Peru ~$0.8–1.8, US ~$5–15 (secondary)."),
    ("incomeTaxRate", "Income tax (RMT, first 15 UIT)", PCT, "10% of profit."),
    ("initialInvestment", "Up-front cash, $", USD, "INDECOPI classes 9 + 42 (~S/ 1,068), domain, accountant set-up."),
    ("discountRateAnnual", "Discount rate, a year", PCT, "Barely matters over 36 months (see Sensitivity)."),
    ("hoursInvested", "Your hours so far", NUM, "Change to your real number."),
    ("hourlyRateUsd", "Your hour, $", USD, "Lima senior developer ~S/ 120."),
    ("conversionMultiplier", "Conversion multiplier", '0.00', "1 = as assumed. Used by price scenarios."),
]


def build_assumptions(wb, ws, a):
    ws["A1"] = "Assumptions"
    ws["A1"].font = TITLE
    ws["A2"] = "Change only yellow cells. Scenario columns B and D copy Base unless a value is typed over them."
    header(ws, 4, ["Demand and conversion", "Pessimistic", "Base", "Optimistic", "Source / note"])
    col = {"pessimistic": "B", "base": "C", "optimistic": "D"}
    for i, (key, label, fmt, note) in enumerate(SCENARIO_ROWS, start=5):
        ws.cell(row=i, column=1, value=label).font = Font(name=FONT)
        for scen, letter in col.items():
            cell = ws[f"{letter}{i}"]
            value = a["scenarios"][scen][key]
            if scen == "base" or value != a["scenarios"]["base"][key]:
                input_cell(cell, value, fmt)
            else:
                cell.value = f"=C{i}"
                cell.font = BLACK
                cell.number_format = fmt
            name(wb, f"{key}_{scen}", f"Assumptions!${letter}${i}")
        ws.cell(row=i, column=5, value=note).font = Font(name=FONT, italic=True)
    start = 5 + len(SCENARIO_ROWS) + 1
    header(ws, start, ["Prices, fees, costs (all scenarios)", "", "Value", "", "Source / note"])
    for i, (key, label, fmt, note) in enumerate(COMMON_ROWS, start=start + 1):
        ws.cell(row=i, column=1, value=label).font = Font(name=FONT)
        input_cell(ws[f"C{i}"], a["common"][key], fmt)
        ws.cell(row=i, column=5, value=note).font = Font(name=FONT, italic=True)
        name(wb, key, f"Assumptions!$C${i}")
    row = start + 1 + len(COMMON_ROWS) + 1
    derived = [
        ("fixedMonthlyUsd", "Fixed compliance cost, $ a month", "=(contadorPen+invoicingPen)/penPerUsd", USD),
        ("monthlyRate", "Monthly discount rate", "=(1+discountRateAnnual)^(1/12)-1", '0.000%'),
    ]
    header(ws, row, ["Derived (formulas)", "", "Value", "", ""])
    for i, (key, label, formula, fmt) in enumerate(derived, start=row + 1):
        ws.cell(row=i, column=1, value=label).font = Font(name=FONT)
        ws[f"C{i}"] = formula
        ws[f"C{i}"].font = BLACK
        ws[f"C{i}"].number_format = fmt
        name(wb, key, f"Assumptions!$C${i}")
    ws.column_dimensions["A"].width = 44
    for letter in "BCD":
        ws.column_dimensions[letter].width = 14
    ws.column_dimensions["E"].width = 90


# ── Unit economics ───────────────────────────────────────────────────────────
def net_formula(gross_usd, tax, fx):
    # Mirrors netPerCharge(): (gross − tax − fee) × (1 − fx)
    return f"=({gross_usd}-({gross_usd}-{gross_usd}/(1+{tax}))-({gross_usd}*feePct+feeFixed))*(1-{fx})"


def build_unit(wb, ws):
    ws["A1"] = "What one charge leaves you"
    ws["A1"].font = TITLE
    ws["A2"] = "Prices include tax. Paddle is the merchant of record: it remits the tax and keeps 5% + $0.50 of the gross."
    header(ws, 4, ["Plan", "Price shown", "Gross in $", "Tax in $", "Paddle fee $", "FX $", "You keep $", "You keep S/", "Share kept"])
    plans = [
        ("Peru, monthly", "priceMonthlyPEN", "PEN", "keptMonthlyPEN"),
        ("Peru, yearly", "priceYearlyPEN", "PEN", "keptYearlyPEN"),
        ("World, monthly", "priceMonthlyUSD", "USD", "keptMonthlyUSD"),
        ("World, yearly", "priceYearlyUSD", "USD", "keptYearlyUSD"),
    ]
    for i, (label, price, cur, kept) in enumerate(plans, start=5):
        tax = "taxPeru" if cur == "PEN" else "taxIntl"
        fx = "fxSpread" if cur == "PEN" else "0"
        ws[f"A{i}"] = label
        ws[f"B{i}"] = f"={price}"
        ws[f"B{i}"].font = GREEN
        ws[f"B{i}"].number_format = PEN if cur == "PEN" else USD
        ws[f"C{i}"] = f"={price}/penPerUsd" if cur == "PEN" else f"={price}"
        ws[f"D{i}"] = f"=C{i}-C{i}/(1+{tax})"
        ws[f"E{i}"] = f"=C{i}*feePct+feeFixed"
        ws[f"F{i}"] = f"=(C{i}-D{i}-E{i})*{fx}"
        ws[f"G{i}"] = net_formula(f"C{i}", tax, fx)
        ws[f"H{i}"] = f"=G{i}*penPerUsd"
        ws[f"I{i}"] = f"=IF(C{i}=0,0,G{i}/C{i})"
        for letter in "CDEFG":
            ws[f"{letter}{i}"].number_format = USD
        ws[f"H{i}"].number_format = PEN
        ws[f"I{i}"].number_format = PCT
        name(wb, kept, f"'Unit economics'!$G${i}")
    ws["A11"] = "Why the soles monthly plan keeps least: the fixed $0.50 is ~13% of a S/ 14.90 charge before IGV, and IGV (18%) comes out of the shown price."
    ws["A11"].font = Font(name=FONT, italic=True)
    ws.column_dimensions["A"].width = 18
    for c in range(2, 10):
        ws.column_dimensions[get_column_letter(c)].width = 14


# ── Cash flow ────────────────────────────────────────────────────────────────
BLOCK_COLS = ["t", "Month", "Sign-ups", "Beta?", "Trial → paid", "Free → paid", "New payers", "New yearly",
              "New monthly", "Yearly charges", "Monthly subs", "Free users", "Active yearly", "Revenue $",
              "Active users", "Infra $", "Ads $", "Costs $", "Pre-tax $", "Net cash $", "Cumulative $",
              "Discounted $", "MRR $", "Paying users", "Paid back?"]


def block_letters(start_col):
    return {k: get_column_letter(start_col + i) for i, k in enumerate(BLOCK_COLS)}


def infra_formula(mau):
    return (f"=IF({mau}<=mau1,cost1,IF({mau}<=mau2,cost1+({mau}-mau1)/(mau2-mau1)*(cost2-cost1),"
            f"IF({mau}<=mau3,cost2+({mau}-mau2)/(mau3-mau2)*(cost3-cost2),cost3*{mau}/mau3)))")


def build_block(ws, scen, start_col, net_m, net_y):
    L = block_letters(start_col)
    s = lambda key: f"{key}_{scen}"  # noqa: E731
    first, last = FIRST_ROW, FIRST_ROW + MONTHS - 1
    header(ws, FIRST_ROW - 1, BLOCK_COLS, start_col)
    for t in range(MONTHS):
        r = FIRST_ROW + t
        p = r - 1
        c = {k: f"{v}{r}" for k, v in L.items()}
        prev = {k: f"{v}{p}" for k, v in L.items()}
        ws[c["t"]] = t
        ws[c["Month"]] = t + 1
        ws[c["Sign-ups"]] = f"={s('signups0')}" if t == 0 else f"={prev['Sign-ups']}*(1+{s('signupGrowth')}*{s('growthFade')}^{prev['t']})"
        ws[c["Beta?"]] = f"=IF({c['t']}<{s('betaMonths')},1,0)"
        ws[c["Trial → paid"]] = f"=IF({c['Beta?']}=1,0,{c['Sign-ups']}*{s('trialStartRate')}*{s('trialToPaid')})"
        prev_free = "0" if t == 0 else prev["Free users"]
        ws[c["Free → paid"]] = f"=IF({c['Beta?']}=1,0,{prev_free}*{s('freemiumMonthly')})"
        ws[c["New payers"]] = f"=({c['Trial → paid']}+{c['Free → paid']})*conversionMultiplier"
        ws[c["New yearly"]] = f"={c['New payers']}*{s('annualShare')}"
        ws[c["New monthly"]] = f"={c['New payers']}-{c['New yearly']}"
        renew = f"+{L['Yearly charges']}{r - 12}*{s('annualRenewal')}" if t >= 12 else ""
        ws[c["Yearly charges"]] = f"={c['New yearly']}{renew}"
        prev_monthly = "0" if t == 0 else prev["Monthly subs"]
        ws[c["Monthly subs"]] = f"={prev_monthly}*(1-{s('monthlyChurn')})+{c['New monthly']}"
        ws[c["Free users"]] = f"={prev_free}*{s('freeRetention')}+{c['Sign-ups']}-{c['Trial → paid']}-{c['Free → paid']}"
        ws[c["Active yearly"]] = f"=SUM({L['Yearly charges']}{max(first, r - 11)}:{c['Yearly charges']})"
        ws[c["Revenue $"]] = f"={c['Monthly subs']}*{net_m}+{c['Yearly charges']}*{net_y}"
        ws[c["Active users"]] = f"={c['Free users']}+{c['Monthly subs']}+{c['Active yearly']}"
        ws[c["Infra $"]] = infra_formula(c["Active users"])
        ws[c["Ads $"]] = f"=IF({c['t']}>=adsStartMonth,{c['Free users']}*adPageviewsPerFreeUser*adRpmUsd/1000,0)"
        ws[c["Costs $"]] = f"=IF({c['Beta?']}=1,{c['Infra $']},{c['Infra $']}+fixedMonthlyUsd+IF({c['Revenue $']}>100,payoutFeeMonthly,0))"
        ws[c["Pre-tax $"]] = f"={c['Revenue $']}+{c['Ads $']}-{c['Costs $']}"
        ws[c["Net cash $"]] = f"=IF({c['Pre-tax $']}>0,{c['Pre-tax $']}*(1-incomeTaxRate),{c['Pre-tax $']})"
        ws[c["Cumulative $"]] = f"=-initialInvestment+{c['Net cash $']}" if t == 0 else f"={prev['Cumulative $']}+{c['Net cash $']}"
        ws[c["Discounted $"]] = f"={c['Net cash $']}/(1+monthlyRate)^({c['t']}+1)"
        ws[c["MRR $"]] = f"={c['Monthly subs']}*{net_m}+{c['Active yearly']}*{net_y}/12"
        ws[c["Paying users"]] = f"={c['Monthly subs']}+{c['Active yearly']}"
        ws[c["Paid back?"]] = f"=IF({c['Cumulative $']}>=0,1,0)"
        for k in BLOCK_COLS:
            cell = ws[c[k]]
            cell.font = Font(name=FONT)
            cell.number_format = USD if k.endswith("$") else NUM1 if k not in ("t", "Month", "Beta?", "Paid back?") else "0"
    return L


def build_cashflow(wb, ws):
    ws["A1"] = "Thirty-six months, three scenarios"
    ws["A1"].font = TITLE
    ws["A2"] = ("Months in the beta charge nobody. Yearly plans are paid up front and renew after 12 months. "
                "Growth fades each month. Same arithmetic as server/billing/financeModel.js.")
    summary_labels = [
        ("Blended kept per monthly charge ($)", USD),
        ("Blended kept per yearly charge ($)", USD),
        ("Net present value over 36 months ($)", USD0),
        ("Cash in over 36 months, after the up-front cost ($)", USD0),
        ("Deepest point below zero ($)", USD0),
        ("Month you have your money back", "0"),
        ("Return on your hours", MULT),
        ("Paying users at month 36", NUM),
        ("MRR at month 36 ($)", USD0),
        ("Active users at month 36", NUM),
    ]
    for i, (label, _) in enumerate(summary_labels, start=4):
        ws.cell(row=i, column=1, value=label).font = BOLD
    width = len(BLOCK_COLS) + 1
    for n, scen in enumerate(["pessimistic", "base", "optimistic"]):
        start_col = 3 + n * width
        col = get_column_letter(start_col + 2)
        ws.cell(row=3, column=start_col + 2, value=scen.upper()).font = BOLD
        net_m_ref = f"${col}$4"
        net_y_ref = f"${col}$5"
        ws[f"{col}4"] = f"={s_(scen, 'peruShare')}*keptMonthlyPEN+(1-{s_(scen, 'peruShare')})*keptMonthlyUSD"
        ws[f"{col}5"] = f"={s_(scen, 'peruShare')}*keptYearlyPEN+(1-{s_(scen, 'peruShare')})*keptYearlyUSD"
        L = build_block(ws, scen, start_col, net_m_ref, net_y_ref)
        rng = lambda k: f"{L[k]}{FIRST_ROW}:{L[k]}{FIRST_ROW + MONTHS - 1}"  # noqa: E731
        last = FIRST_ROW + MONTHS - 1
        ws[f"{col}6"] = f"=-initialInvestment+SUM({rng('Discounted $')})"
        ws[f"{col}7"] = f"={L['Cumulative $']}{last}"
        ws[f"{col}8"] = f"=MIN(0,-initialInvestment,MIN({rng('Cumulative $')}))"
        ws[f"{col}9"] = f"=IFERROR(INDEX({rng('Month')},MATCH(1,{rng('Paid back?')},0)),\"not within 36 months\")"
        ws[f"{col}10"] = f"=IF(hoursInvested*hourlyRateUsd=0,0,{col}7/(hoursInvested*hourlyRateUsd))"
        ws[f"{col}11"] = f"={L['Paying users']}{last}"
        ws[f"{col}12"] = f"={L['MRR $']}{last}"
        ws[f"{col}13"] = f"={L['Active users']}{last}"
        for i, (_, fmt) in enumerate(summary_labels, start=4):
            ws[f"{col}{i}"].number_format = fmt
            ws[f"{col}{i}"].font = GREEN if i in (4, 5) else BLACK
        name(wb, f"netcash_{scen}", f"'Cash flow'!${L['Net cash $']}${FIRST_ROW}:${L['Net cash $']}${last}")
        name(wb, f"tplus1_{scen}", f"'Cash flow'!${L['Month']}${FIRST_ROW}:${L['Month']}${last}")
        name(wb, f"npv_{scen}", f"'Cash flow'!${col}$6")
    ws.column_dimensions["A"].width = 46
    ws.freeze_panes = ws.cell(row=FIRST_ROW, column=3)


def s_(scen, key):
    return f"{key}_{scen}"


# ── Break-even ───────────────────────────────────────────────────────────────
def build_breakeven(ws):
    ws["A1"] = "Break-even"
    ws["A1"].font = TITLE
    ws["A2"] = "Base scenario mix. Fixed costs start when the RUC does; before that only hosting costs anything."
    ws["A4"] = "Kept per paying user per month ($)"
    ws["B4"] = "=annualShare_base*(peruShare_base*keptYearlyPEN+(1-peruShare_base)*keptYearlyUSD)/12+(1-annualShare_base)*(peruShare_base*keptMonthlyPEN+(1-peruShare_base)*keptMonthlyUSD)"
    ws["B4"].number_format = USD
    ws["B4"].font = GREEN
    header(ws, 6, ["Active users", "Hosting $", "Compliance $", "Payout fee $", "Total fixed $", "Paying users needed"])
    for i, mau in enumerate([1000, 10000, 50000], start=7):
        input_cell(ws[f"A{i}"], mau, NUM)
        ws[f"B{i}"] = infra_formula(f"A{i}")
        ws[f"C{i}"] = "=fixedMonthlyUsd"
        ws[f"D{i}"] = "=payoutFeeMonthly"
        ws[f"E{i}"] = f"=B{i}+C{i}+D{i}"
        ws[f"F{i}"] = f"=ROUNDUP(E{i}/$B$4,0)"
        for letter in "BCDE":
            ws[f"{letter}{i}"].number_format = USD
        ws[f"F{i}"].number_format = NUM
    ws["A12"] = "The contador is the cost structure"
    ws["A12"].font = BOLD
    header(ws, 13, ["Contador S/ a month", "Fixed $ a month (2,000 active users)", "Paying users needed"])
    for i, soles in enumerate([0, 150, 250, 350, 500], start=14):
        input_cell(ws[f"A{i}"], soles, PEN)
        ws[f"B{i}"] = f"=(A{i}+invoicingPen)/penPerUsd+{infra_formula('2000')[1:]}+payoutFeeMonthly"
        ws[f"C{i}"] = f"=ROUNDUP(B{i}/$B$4,0)"
        ws[f"B{i}"].number_format = USD
        ws[f"C{i}"].number_format = NUM
    ws["A20"] = "Income-tax headroom: RMT is 10% up to 15 UIT = S/ 82,500 of annual profit (UIT 2026 = S/ 5,500)."
    ws["A21"] = "Paying users where yearly profit reaches 15 UIT (approx.)"
    ws["B21"] = "=ROUNDUP((82500/penPerUsd/12+fixedMonthlyUsd)/$B$4,0)"
    ws["B21"].number_format = NUM
    for col, width in zip("ABCDEF", [34, 22, 18, 16, 16, 20]):
        ws.column_dimensions[col].width = width


# ── Price equilibrium ────────────────────────────────────────────────────────
def build_price(ws):
    ws["A1"] = "Price equilibrium — soles monthly plan"
    ws["A1"].font = TITLE
    ws["A2"] = ("Linear demand: a 1% price rise loses β% of buyers, indexed to 100 buyers at the reference price. "
                "Each cell is buyers × what you keep, so only the shape of a column means anything.")
    ws["A4"] = "Reference price (S/)"
    ws["B4"] = "=priceMonthlyPEN"
    ws["B4"].font = GREEN
    ws["B4"].number_format = PEN
    header(ws, 6, ["Price S/", "You keep S/", "β = 0.8", "β = 1.2", "β = 1.6", "β = 2.0"])
    betas = [0.8, 1.2, 1.6, 2.0]
    for j, beta in enumerate(betas):
        input_cell(ws.cell(row=5, column=3 + j), beta, '0.0')
    ws["B5"] = "β (edit) →"
    for i, price in enumerate([9.90, 12.90, 14.90, 16.90, 19.90, 22.90, 24.90, 29.90], start=7):
        input_cell(ws[f"A{i}"], price, PEN)
        g = f"(A{i}/penPerUsd)"
        ws[f"B{i}"] = f"=({g}-({g}-{g}/(1+taxPeru))-({g}*feePct+feeFixed))*(1-fxSpread)*penPerUsd"
        ws[f"B{i}"].number_format = PEN
        for j in range(len(betas)):
            beta_cell = f"{get_column_letter(3 + j)}$5"
            cell = ws.cell(row=i, column=3 + j)
            cell.value = f"=MAX(0,100*(1-{beta_cell}*(A{i}/$B$4-1)))*B{i}"
            cell.number_format = NUM
    ws["A16"] = ("Read down a column: the peak is the best price for that elasticity. At β 1.2–1.6 the peak is S/ 12.90–14.90; "
                 "at β 0.8 it moves to S/ 16.90. Test before moving: a price survey of beta testers, not a price A/B test.")
    ws["A16"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A16:F17")
    for col in "ABCDEF":
        ws.column_dimensions[col].width = 16


# ── Sensitivity + engine ─────────────────────────────────────────────────────
ENGINE_COLS = ["Trial → paid", "Free → paid", "Free users", "New yearly", "New monthly", "Yearly charges",
               "Monthly subs", "Active yearly", "Revenue $", "Active users", "Infra $", "Ads $", "Costs $",
               "Net cash $", "Discounted $"]


def build_engine(wb, ws, conversions, churns):
    ws["A1"] = "Sensitivity engine — base scenario re-run for each trial-to-paid × churn pair"
    ws["A1"].font = TITLE
    header(ws, 3, ["t", "Sign-ups", "Beta?"])
    first = 4
    for t in range(MONTHS):
        r = first + t
        ws[f"A{r}"] = t
        ws[f"B{r}"] = "=signups0_base" if t == 0 else f"=B{r - 1}*(1+signupGrowth_base*growthFade_base^A{r - 1})"
        ws[f"C{r}"] = f"=IF(A{r}<betaMonths_base,1,0)"
    net_m = "(peruShare_base*keptMonthlyPEN+(1-peruShare_base)*keptMonthlyUSD)"
    net_y = "(peruShare_base*keptYearlyPEN+(1-peruShare_base)*keptYearlyUSD)"
    npv_cells = {}
    col = 5
    for ci, conv in enumerate(conversions):
        for hi, churn in enumerate(churns):
            L = {k: get_column_letter(col + i) for i, k in enumerate(ENGINE_COLS)}
            ws.cell(row=2, column=col, value=f"conv {conv:.0%} · churn {churn:.0%}").font = BOLD
            header(ws, 3, ENGINE_COLS, col)
            conv_ref = f"Sensitivity!$A${9 + ci}"
            churn_ref = f"Sensitivity!${get_column_letter(2 + hi)}$8"
            for t in range(MONTHS):
                r = first + t
                c = {k: f"{v}{r}" for k, v in L.items()}
                pv = {k: f"{v}{r - 1}" for k, v in L.items()}
                prev_free = "0" if t == 0 else pv["Free users"]
                prev_monthly = "0" if t == 0 else pv["Monthly subs"]
                ws[c["Trial → paid"]] = f"=IF($C{r}=1,0,$B{r}*trialStartRate_base*{conv_ref})"
                ws[c["Free → paid"]] = f"=IF($C{r}=1,0,{prev_free}*freemiumMonthly_base)"
                ws[c["Free users"]] = f"={prev_free}*freeRetention_base+$B{r}-{c['Trial → paid']}-{c['Free → paid']}"
                payers = f"({c['Trial → paid']}+{c['Free → paid']})*conversionMultiplier"
                ws[c["New yearly"]] = f"={payers}*annualShare_base"
                ws[c["New monthly"]] = f"={payers}-{c['New yearly']}"
                renew = f"+{L['Yearly charges']}{r - 12}*annualRenewal_base" if t >= 12 else ""
                ws[c["Yearly charges"]] = f"={c['New yearly']}{renew}"
                ws[c["Monthly subs"]] = f"={prev_monthly}*(1-{churn_ref})+{c['New monthly']}"
                ws[c["Active yearly"]] = f"=SUM({L['Yearly charges']}{max(first, r - 11)}:{c['Yearly charges']})"
                ws[c["Revenue $"]] = f"={c['Monthly subs']}*{net_m}+{c['Yearly charges']}*{net_y}"
                ws[c["Active users"]] = f"={c['Free users']}+{c['Monthly subs']}+{c['Active yearly']}"
                ws[c["Infra $"]] = infra_formula(c["Active users"])
                ws[c["Ads $"]] = f"=IF($A{r}>=adsStartMonth,{c['Free users']}*adPageviewsPerFreeUser*adRpmUsd/1000,0)"
                ws[c["Costs $"]] = f"=IF($C{r}=1,{c['Infra $']},{c['Infra $']}+fixedMonthlyUsd+IF({c['Revenue $']}>100,payoutFeeMonthly,0))"
                pre = f"({c['Revenue $']}+{c['Ads $']}-{c['Costs $']})"
                ws[c["Net cash $"]] = f"=IF({pre}>0,{pre}*(1-incomeTaxRate),{pre})"
                ws[c["Discounted $"]] = f"={c['Net cash $']}/(1+monthlyRate)^($A{r}+1)"
            last = first + MONTHS - 1
            npv_cells[(ci, hi)] = f"=-initialInvestment+SUM('Sensitivity engine'!{L['Discounted $']}{first}:{L['Discounted $']}{last})"
            col += len(ENGINE_COLS) + 1
    return npv_cells


def build_sensitivity(ws, npv_cells, conversions, churns):
    ws["A1"] = "What moves the answer"
    ws["A1"].font = TITLE
    # Conversion × churn grid (rows 8..): churn across, trial-to-paid down.
    ws["A7"] = "NPV over 36 months, base scenario: trial-to-paid down the side, monthly churn across the top"
    ws["A7"].font = BOLD
    ws["A8"] = "Trial → paid ↓ / churn →"
    for hi, churn in enumerate(churns):
        input_cell(ws.cell(row=8, column=2 + hi), churn, PCT)
    for ci, conv in enumerate(conversions):
        input_cell(ws.cell(row=9 + ci, column=1), conv, PCT)
        for hi in range(len(churns)):
            cell = ws.cell(row=9 + ci, column=2 + hi, value=npv_cells[(ci, hi)])
            cell.number_format = USD0
            cell.font = GREEN
    ws["A16"] = "NPV against the discount rate (base scenario)"
    ws["A16"].font = BOLD
    header(ws, 17, ["Discount rate a year", "NPV $"])
    for i, rate in enumerate([0.08, 0.12, 0.15, 0.20, 0.30, 0.40]):
        r = 18 + i
        input_cell(ws[f"A{r}"], rate, PCT)
        ws[f"B{r}"] = f"=-initialInvestment+SUMPRODUCT(netcash_base/((1+A{r})^(tplus1_base/12)))"
        ws[f"B{r}"].number_format = USD0
        ws[f"B{r}"].font = GREEN
    ws["A25"] = "Quadrupling the discount rate moves NPV far less than conversion does; don't spend thought on the rate."
    ws["A25"].font = Font(name=FONT, italic=True)
    for i in range(1, 8):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.column_dimensions["A"].width = 30


def main(out):
    a = assumptions()
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    sheets = {n: wb.create_sheet(n) for n in ["Assumptions", "Unit economics", "Break-even", "Price equilibrium",
                                               "Cash flow", "Sensitivity", "Sensitivity engine"]}
    build_readme(readme)
    build_assumptions(wb, sheets["Assumptions"], a)
    build_unit(wb, sheets["Unit economics"])
    build_breakeven(sheets["Break-even"])
    build_price(sheets["Price equilibrium"])
    build_cashflow(wb, sheets["Cash flow"])
    conversions = [0.06, 0.08, 0.10, 0.14, 0.18]
    churns = [0.05, 0.07, 0.09, 0.12, 0.15]
    npv_cells = build_engine(wb, sheets["Sensitivity engine"], conversions, churns)
    build_sensitivity(sheets["Sensitivity"], npv_cells, conversions, churns)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.font is None or cell.font.name != FONT:
                    cell.font = Font(name=FONT, bold=cell.font.bold if cell.font else False,
                                     color=cell.font.color if cell.font else None)
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "gylio-financial-model.xlsx")
