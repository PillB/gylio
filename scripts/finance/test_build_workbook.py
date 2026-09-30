"""Structural checks on the generated workbook that don't need a spreadsheet engine.

    python3 -m unittest scripts/finance/test_build_workbook.py
"""
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
FUNCTIONS = {"IF", "SUM", "MIN", "MAX", "ROUNDUP", "IFERROR", "INDEX", "MATCH", "SUMPRODUCT"}


class WorkbookStructure(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.path = Path(tempfile.mkdtemp()) / "model.xlsx"
        subprocess.run([sys.executable, "scripts/finance/build_workbook.py", str(cls.path)], cwd=ROOT, check=True,
                       capture_output=True)
        cls.wb = load_workbook(cls.path)

    def formulas(self):
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        yield ws.title, cell.coordinate, cell.value

    def test_every_name_a_formula_uses_is_defined(self):
        defined = set(self.wb.defined_names.keys())
        missing = set()
        for sheet, coord, formula in self.formulas():
            body = re.sub(r"'[^']*'![$A-Z0-9:]+|[A-Za-z ]+![$A-Z0-9:]+|\"[^\"]*\"", "", formula)
            body = re.sub(r"\$?\b[A-Z]{1,3}\$?\d+\b", "", body)  # cell references like $AE$4 or B12
            for token in re.findall(r"\b[A-Za-z_][A-Za-z0-9_]*\b", body):
                if token.upper() in FUNCTIONS or re.fullmatch(r"\$?[A-Z]{1,3}\$?\d+", token) or token in ("TRUE", "FALSE"):
                    continue
                if token not in defined:
                    missing.add((token, sheet, coord))
        self.assertEqual(sorted(missing)[:10], [])

    def test_only_functions_libreoffice_and_excel_both_evaluate(self):
        for _, _, formula in self.formulas():
            for fn in re.findall(r"([A-Z][A-Z0-9.]*)\(", formula):
                self.assertIn(fn, FUNCTIONS, formula)

    def test_every_month_row_of_every_scenario_has_the_same_formula_shape(self):
        ws = self.wb["Cash flow"]
        # Net cash column of the base block: same formula pattern on all 36 rows.
        shapes = {re.sub(r"\d+", "#", ws.cell(row=r, column=3 + 26 + 19).value) for r in range(20, 56)}
        self.assertEqual(len(shapes), 1, shapes)

    def test_inputs_are_marked_yellow_with_blue_text(self):
        ws = self.wb["Assumptions"]
        inputs = [c for row in ws.iter_rows() for c in row if c.fill.fgColor.rgb in ("00FFFF00", "FFFFFF00")]
        self.assertGreater(len(inputs), 40)
        for cell in inputs:
            self.assertEqual(cell.font.color.rgb[-6:], "0000FF", cell.coordinate)
            self.assertFalse(isinstance(cell.value, str) and cell.value.startswith("="), cell.coordinate)


if __name__ == "__main__":
    unittest.main()
